"""Bundeslebensmittelschlüssel (BLS) 4.0 – Import der deutschen Nährstoffdatenbank.

Quelle: Max Rubner-Institut, Bundeslebensmittelschlüssel (BLS) Version 4.0, https://www.blsdb.de
Lizenz: CC BY 4.0. Enthält ca. 7.140 generische Lebensmittel und Gerichte (z. B. Döner Kebab, Gyros, Lahmacun).

Aufbau der Datei: Spalten 1–3 = BLS-Code, deutscher Name, englischer Name; danach je Nährstoff drei Spalten
(Wert, Datenquelle, Referenz). Wertspalten beginnen mit dem EuroFIR-Code, z. B. „PROT625 Protein … [g/100g]“.
"""

import csv
import io
import logging
import re
from collections.abc import Iterator
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

import httpx
from sqlmodel import select
from sqlmodel.ext.asyncio.session import AsyncSession

from app.core.config import settings
from app.models import Food
from app.services.app_settings import get_app_setting, set_app_setting

log = logging.getLogger(__name__)

ATTRIBUTION = "Bundeslebensmittelschlüssel (BLS) 4.0, Max Rubner-Institut, CC BY 4.0"
DOWNLOAD_PAGE = "https://www.blsdb.de/download"
FALLBACK_URLS = [
    "https://www.blsdb.de/assets/uploads/BLS_4_0_Daten_2025_DE.xlsx",
    "https://blsdb.de/assets/uploads/BLS_4_0_Daten_2025_DE.xlsx",
]

# Zielfeld -> mögliche Komponenten-Codes (erste passende Spalte gewinnt) und Zieleinheit
FIELDS: dict[str, tuple[tuple[str, ...], str]] = {
    "kcal": (("ENERCC",), "kcal"),
    "protein": (("PROT625", "PROT"), "g"),
    "fat": (("FAT",), "g"),
    "carbs": (("CHO",), "g"),
    "fiber": (("FIBT",), "g"),
    "sugar": (("SUGAR",), "g"),
    "sat_fat": (("FASAT",), "g"),
    "salt": (("NACL",), "g"),
    "_sodium": (("NA",), "mg"),
    "_kj": (("ENERCJ",), "kJ"),
}
MICROS: dict[str, tuple[str, str]] = {
    "vitamin_c_mg": ("VITC", "mg"), "vitamin_d_ug": ("VITD", "µg"), "vitamin_b12_ug": ("VITB12", "µg"),
    "calcium_mg": ("CA", "mg"), "iron_mg": ("FE", "mg"), "magnesium_mg": ("MG", "mg"), "potassium_mg": ("K", "mg"),
    "zinc_mg": ("ZN", "mg"), "sodium_mg": ("NA", "mg"),
}
UNIT_FACTOR = {("g", "g"): 1, ("mg", "g"): 0.001, ("µg", "g"): 1e-6, ("g", "mg"): 1000, ("mg", "mg"): 1, ("µg", "mg"): 0.001,
               ("g", "µg"): 1e6, ("mg", "µg"): 1000, ("µg", "µg"): 1, ("kcal", "kcal"): 1, ("kj", "kj"): 1}

# Hauptgruppen nach erstem Buchstaben des BLS-Codes
GROUPS = {
    "B": "Brot & Kleingebäck", "C": "Getreide", "D": "Backwaren & Kuchen", "E": "Eier & Teigwaren", "F": "Obst",
    "G": "Gemüse", "H": "Hülsenfrüchte & Nüsse", "K": "Kartoffeln & Pilze", "M": "Milch & Milchprodukte",
    "N": "Getränke", "P": "Alkoholische Getränke", "Q": "Fette & Öle", "R": "Gewürze & Würzmittel",
    "S": "Süßwaren & Zucker", "T": "Fisch & Meeresfrüchte", "U": "Fleisch", "V": "Wild, Geflügel & Innereien",
    "W": "Wurst & Aufschnitt", "X": "Gerichte & Snacks", "Y": "Gerichte & Snacks",
}

# Typische Portionen (g) nach Stichwort – BLS enthält keine Portionsgrößen
PORTIONS: list[tuple[str, float, str]] = [
    ("döner", 380, "1 Döner"), ("dürüm", 380, "1 Dürüm"), ("lahmacun", 280, "1 Lahmacun"), ("gyrosburger", 250, "1 Burger"),
    ("gyros", 200, "1 Portion"), ("pizzabaguette", 125, "1 Stück"), ("pizza", 380, "1 Pizza"), ("burger", 220, "1 Burger"),
    ("hamburger", 220, "1 Burger"), ("cheeseburger", 230, "1 Burger"), ("pommes", 150, "1 Portion"), ("currywurst", 250, "1 Portion"),
    ("bratwurst", 120, "1 Wurst"), ("brötchen", 55, "1 Brötchen"), ("brezel", 80, "1 Brezel"), ("croissant", 60, "1 Stück"),
    ("schnitzel", 180, "1 Schnitzel"), ("frikadelle", 100, "1 Stück"), ("ei ", 60, "1 Ei"), ("apfel", 150, "1 Apfel"),
    ("banane", 120, "1 Banane"), ("joghurt", 150, "1 Becher"), ("müsli", 60, "1 Portion"), ("haferflocken", 50, "1 Portion"),
    ("reis", 180, "1 Portion gekocht"), ("teigwaren", 200, "1 Portion gekocht"), ("nudeln", 200, "1 Portion gekocht"),
    ("kartoffel", 200, "1 Portion"), ("salat", 150, "1 Schale"), ("suppe", 300, "1 Teller"), ("eintopf", 350, "1 Teller"),
    ("kuchen", 120, "1 Stück"), ("torte", 130, "1 Stück"), ("schokolade", 25, "1 Riegel/Stück"), ("bier", 500, "0,5 l"),
    ("wein", 200, "1 Glas"), ("milch", 250, "1 Glas"), ("saft", 250, "1 Glas"),
]


class BLSImportError(Exception):
    pass


def _num(v: Any) -> float | None:
    if v is None:
        return None
    if isinstance(v, int | float):
        return float(v)
    s = str(v).strip().replace(",", ".")
    if not s or s in ("-", "–", "n.a.", "NA"):
        return None
    if s.startswith("<") or s.lower() in ("tr", "trace", "spuren"):
        return 0.0
    try:
        return float(s)
    except ValueError:
        return None


def _header_code_unit(h: str) -> tuple[str, str]:
    h = str(h or "").strip()
    code = h.split(" ", 1)[0].upper()
    m = re.search(r"\[([^\]/]+)/100\s*g\]", h)
    unit = (m.group(1).strip() if m else "").replace("ug", "µg").replace("mcg", "µg")
    return code, unit.lower() if unit.lower() in ("kcal", "kj") else unit


def map_columns(headers: list[str]) -> dict[str, tuple[int, str]]:
    """Ordnet Zielfelder den Wertspalten zu (Index, Einheit)."""
    found: dict[str, tuple[int, str]] = {}
    codes = [_header_code_unit(h) for h in headers]
    for field, (cands, _unit) in FIELDS.items():
        for cand in cands:
            idx = next((i for i, (c, u) in enumerate(codes) if c == cand and u), None)
            if idx is not None:
                found[field] = (idx, codes[idx][1])
                break
    for field, (cand, _unit) in MICROS.items():
        idx = next((i for i, (c, u) in enumerate(codes) if c == cand and u), None)
        if idx is not None:
            found["micro:" + field] = (idx, codes[idx][1])
    missing = [f for f in ("protein", "fat", "carbs") if f not in found] + ([] if ("kcal" in found or "_kj" in found) else ["kcal"])
    if missing:
        raise BLSImportError(f"Spalten nicht gefunden: {', '.join(missing)} – ist das die BLS-4.0-Datei?")
    return found


def _convert(value: float | None, unit: str, target: str) -> float | None:
    if value is None:
        return None
    return value * UNIT_FACTOR.get((unit, target), 1)


def portion_for(name: str) -> tuple[float | None, str]:
    low = name.lower() + " "
    for key, grams, label in PORTIONS:
        if key in low:
            return grams, label
    return None, ""


def rows_to_foods(headers: list[str], rows: Iterator[list[Any]]) -> Iterator[dict[str, Any]]:
    cols = map_columns(headers)
    for row in rows:
        if not row or not row[0] or not row[1]:
            continue
        code = str(row[0]).strip()
        if not re.match(r"^[A-Z][0-9A-Z]{5,7}$", code):
            continue

        def get(field: str, target: str) -> float | None:
            if field not in cols:
                return None
            idx, unit = cols[field]
            return _convert(_num(row[idx]) if idx < len(row) else None, unit, target)

        kcal = get("kcal", "kcal")
        if kcal is None and (kj := get("_kj", "kj")) is not None:
            kcal = kj / 4.184
        salt = get("salt", "g")
        if salt is None and (na := get("_sodium", "mg")) is not None:
            salt = na * 2.5 / 1000
        micros = {}
        for f, (_c, u) in MICROS.items():
            v = get("micro:" + f, u)
            if v:
                micros[f] = round(v, 3)
        name = str(row[1]).strip()
        serving, label = portion_for(name)
        yield {
            "external_id": f"bls:{code}",
            "name": name[:200],
            "brand": "",
            "category": GROUPS.get(code[0], ""),
            "source": "bls",
            "kcal": round(kcal or 0, 1),
            "protein": round(get("protein", "g") or 0, 2),
            "carbs": round(get("carbs", "g") or 0, 2),
            "fat": round(get("fat", "g") or 0, 2),
            "fiber": round(get("fiber", "g") or 0, 2),
            "sugar": round(get("sugar", "g") or 0, 2),
            "sat_fat": round(get("sat_fat", "g") or 0, 2),
            "salt": round(salt or 0, 3),
            "micros": micros,
            "serving_g": serving,
            "serving_label": label,
            "visibility": "public",
        }


def read_file(path: Path) -> Iterator[dict[str, Any]]:
    data = path.read_bytes()
    if data[:2] == b"PK":  # xlsx
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        ws = wb.worksheets[0]
        it = ws.iter_rows(values_only=True)
        headers = [str(h or "") for h in next(it)]
        yield from rows_to_foods(headers, (list(r) for r in it))
        wb.close()
    else:
        text = data.decode("utf-8-sig", errors="replace")
        dialect = csv.Sniffer().sniff(text.splitlines()[0], delimiters=",;\t")
        reader = csv.reader(io.StringIO(text), dialect)
        headers = next(reader)
        yield from rows_to_foods(headers, reader)


async def download(target: Path, url: str | None = None) -> Path:
    """Lädt die BLS-Excel-Datei. Sucht den Link auf der Download-Seite, sonst bekannte URLs."""
    target.parent.mkdir(parents=True, exist_ok=True)
    candidates = [url] if url else []
    async with httpx.AsyncClient(timeout=120, follow_redirects=True, headers={"User-Agent": settings.off_user_agent}) as c:
        if not url:
            try:
                page = await c.get(DOWNLOAD_PAGE)
                for href in re.findall(r'href="([^"]+\.xlsx[^"]*)"', page.text):
                    if "daten" in href.lower() and ("_de" in href.lower() or "DE" in href):
                        candidates.append(str(page.url.join(href)))
            except httpx.HTTPError as e:
                log.warning("BLS-Downloadseite nicht erreichbar: %s", e)
            candidates += FALLBACK_URLS
        last = None
        for u in dict.fromkeys(candidates):
            try:
                r = await c.get(u)
                if r.status_code == 200 and r.content[:2] == b"PK":
                    target.write_bytes(r.content)
                    return target
                last = f"{u}: HTTP {r.status_code}"
            except httpx.HTTPError as e:
                last = f"{u}: {e}"
    raise BLSImportError(f"Download fehlgeschlagen ({last}). Datei unter blsdb.de herunterladen und im Admin-Panel hochladen.")


async def import_foods(db: AsyncSession, foods: Iterator[dict[str, Any]], source_name: str = "BLS 4.0") -> dict[str, Any]:
    existing = {f.external_id: f for f in (await db.exec(select(Food).where(Food.source == "bls"))).all()}
    created = updated = 0
    now = datetime.now(UTC)
    for data in foods:
        f = existing.get(data["external_id"])
        if f:
            for k, v in data.items():
                setattr(f, k, v)
            updated += 1
        else:
            f = Food(owner_id=None, **data)
            created += 1
        f.fetched_at = now
        db.add(f)
        if (created + updated) % 500 == 0:
            await db.flush()
    if not created and not updated:
        raise BLSImportError("Keine Lebensmittel in der Datei gefunden")
    await db.commit()
    status = {"count": created + updated, "created": created, "updated": updated, "imported_at": now.isoformat(),
              "source": source_name, "attribution": ATTRIBUTION, "error": None}
    await set_app_setting(db, "food_db_bls", status)
    log.info("BLS importiert: %d neu, %d aktualisiert", created, updated)
    return status


async def import_bls(db: AsyncSession, path: Path | None = None, url: str | None = None) -> dict[str, Any]:
    try:
        if path is None:
            path = await download(settings.data_dir / "food-db" / "BLS_4_0_Daten_DE.xlsx", url)
        return await import_foods(db, read_file(path))
    except Exception as e:
        status = {**(await get_app_setting(db, "food_db_bls", {})), "error": str(e)[:500],
                  "last_attempt": datetime.now(UTC).isoformat()}
        await set_app_setting(db, "food_db_bls", status)
        raise


async def bls_status(db: AsyncSession) -> dict[str, Any]:
    return {"attribution": ATTRIBUTION, **(await get_app_setting(db, "food_db_bls", {}))}
